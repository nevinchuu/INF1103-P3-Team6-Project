import io
import os
import re
import threading
import uuid

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader

RESUME_FOLDER = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "resumes")
ALLOWED_EXTENSIONS = [".pdf", ".docx"]
MIN_RESUME_CHARS = 100
MAX_SALARY = 50000
MAX_EXPERIENCE_YEARS = 60
LINE_WIDTH = 78

ANY = "Any"
JOB_TYPES = [ANY, "Full Time", "Part Time", "Contract", "Internship/Attachment"]
WORK_ARRANGEMENTS = [ANY, "Onsite", "Hybrid", "Remote"]
EDUCATION_LEVELS = ["None", "Secondary", "ITE/Nitec", "A-Level", "Diploma", "Bachelor's", "Master's", "Doctorate"]

MENU_SEARCH = "Find jobs for my resume"
MENU_RESULTS = "Show my last results"
MENU_DETAILS = "View one job in detail"
MENU_EXIT = "Exit"
MENU_OPTIONS = [MENU_SEARCH, MENU_RESULTS, MENU_DETAILS, MENU_EXIT]
MENU_REFRESH = "Refresh this list"
MENU_CANCEL = "Cancel"

# Short words kept in capitals when an all-caps title is tidied, e.g. "IT SUPPORT" -> "IT Support"
ACRONYMS = {"IT", "HR", "AI", "UX", "UI", "QA", "PR", "F&B", "SQL", "AWS", "SAP", "CRM", "ERP", "IOS", "PHP"}


def display_message(text):
    """Print a normal message."""
    print(text)


def display_error(text):
    """Print an error message."""
    print(f"[!] {text}")


def display_header(title):
    """Print a section title."""
    print()
    print("=" * LINE_WIDTH)
    print(title.center(LINE_WIDTH))
    print("=" * LINE_WIDTH)


def prompt_text(prompt, allow_blank=False):
    """Ask for text until something is typed (or once, if allow_blank)."""
    while True:
        answer = input(f"{prompt}: ").strip()
        if answer or allow_blank:
            return answer
        display_error("This cannot be left blank.")


def prompt_number(prompt, minimum, maximum, default=None):
    """Ask for a whole number within a range."""
    hint = f"{minimum}-{maximum}"
    if default is not None:
        hint += f", Enter for {default}"
    while True:
        answer = input(f"{prompt} ({hint}): ").strip()
        if answer == "" and default is not None:
            return default
        answer = answer.replace(",", "").replace("$", "")
        try:
            value = int(answer)
        except ValueError:
            display_error(f"Please enter a whole number between {minimum} and {maximum}.")
            continue
        if value < minimum or value > maximum:
            display_error(f"Please enter a number between {minimum} and {maximum}.")
            continue
        return value


def prompt_choice(prompt, options):
    """Show numbered options and return the one picked."""
    print()
    print(prompt)
    for number, option in enumerate(options, start=1):
        print(f"  {number}. {option}")
    choice = prompt_number("Choose", 1, len(options))
    return options[choice - 1]


def prompt_yes_no(prompt):
    """Ask a yes/no question."""
    while True:
        answer = input(f"{prompt} (y/n): ").strip().lower()
        if answer in ("y", "yes"):
            return True
        if answer in ("n", "no"):
            return False
        display_error("Please type y or n.")


def check_resume_file(path):
    """Return what is wrong with a file, or an empty string."""
    if not os.path.exists(path):
        return f"File not found: {path}"
    if not os.path.isfile(path):
        return "That is a folder, not a file."
    extension = os.path.splitext(path)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        return f"'{extension}' files are not supported. Please use a PDF or Word (.docx) file."
    if os.path.getsize(path) == 0:
        return "That file is empty."
    return ""


def read_pdf(path):
    """Return the text of a PDF file."""
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def read_docx(path):
    """Return the text of a Word file."""
    document = Document(path)
    lines = []
    for paragraph in document.paragraphs:
        lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return "\n".join(lines)


def read_resume(path):
    """Return the resume text, or an empty string if unreadable."""
    try:
        if path.lower().endswith(".pdf"):
            text = read_pdf(path)
        else:
            text = read_docx(path)
    except Exception as error:
        display_error(f"Could not open the file ({type(error).__name__}).")
        return ""
    return text.strip()


def load_resume(path):
    """Check a resume file and read it. Return (text, problem); problem is empty if it worked."""
    problem = check_resume_file(path)
    if problem:
        return "", problem
    text = read_resume(path)
    if len(text) < MIN_RESUME_CHARS:
        return "", ("Could not read enough text from that file. It may be a scanned image, "
                    "password-protected or corrupted. Please try another file.")
    return text, ""


def write_resume_docx(resume):
    """Return a tailored resume (ai_manager.tailor_resume) as the bytes of a Word file.
    A plain one-column layout, which applicant tracking systems read reliably."""
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(2)

    def heading(text):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)

    if resume["name"]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(resume["name"])
        run.bold = True
        run.font.size = Pt(18)
    if resume["contact"]:
        paragraph = document.add_paragraph(resume["contact"])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if resume["summary"]:
        heading("Summary")
        document.add_paragraph(resume["summary"])
    if resume["skills"]:
        heading("Skills")
        document.add_paragraph(" · ".join(resume["skills"]))

    for section in resume["sections"]:
        heading(section["heading"])
        for entry in section["entries"]:
            # "Title, Organisation" on the left, dates on the right (a right-aligned tab stop)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
            paragraph.add_run(entry["title"]).bold = True
            if entry["organisation"] and entry["organisation"] != entry["title"]:
                paragraph.add_run(f", {entry['organisation']}")
            if entry["dates"]:
                paragraph.add_run(f"\t{entry['dates']}").italic = True
            for bullet in entry["bullets"]:
                document.add_paragraph(bullet, style="List Bullet")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def list_resumes():
    """Return the PDF and Word files in the resumes folder."""
    try:
        os.makedirs(RESUME_FOLDER, exist_ok=True)
        all_names = sorted(os.listdir(RESUME_FOLDER))
    except OSError:
        return []
    names = []
    for name in all_names:
        extension = os.path.splitext(name)[1].lower()
        if extension in ALLOWED_EXTENSIONS and not name.startswith("~$"):
            names.append(name)
    return names


def unused_resume_path(filename):
    """Path in the resumes folder that won't replace a saved resume:
    resume.pdf, or resume_2.pdf, resume_3.pdf... if that name is taken."""
    stem, extension = os.path.splitext(filename)
    path = os.path.join(RESUME_FOLDER, filename)
    number = 2
    while os.path.exists(path):
        path = os.path.join(RESUME_FOLDER, f"{stem}_{number}{extension}")
        number += 1
    return path


def resume_path_from_form(uploaded_name, chosen_name):
    """Where the web form's resume is. Return (path, problem).
    uploaded_name: a safe file name for an upload (the caller saves the file to path), or "" for none.
    chosen_name: the saved resume picked from the list, used when nothing was uploaded."""
    if uploaded_name:
        if os.path.splitext(uploaded_name)[1].lower() not in ALLOWED_EXTENSIONS:
            return "", "Please upload a PDF or Word (.docx) file."
        os.makedirs(RESUME_FOLDER, exist_ok=True)
        return unused_resume_path(uploaded_name), ""
    if chosen_name not in list_resumes():
        return "", "Please choose a resume or upload one."
    return os.path.join(RESUME_FOLDER, chosen_name), ""


def prompt_resume():
    """Let the user pick a resume and return its path and text."""
    while True:
        names = list_resumes()
        if not names:
            display_error("No PDF or Word (.docx) files found. Put your resume in this folder:")
            display_message(f"    {RESUME_FOLDER}")
            answer = input("Press Enter to check again, or q to cancel: ").strip().lower()
            if answer == "q":
                return None
            continue

        choice = prompt_choice("Choose your resume", names + [MENU_REFRESH, MENU_CANCEL])
        if choice == MENU_CANCEL:
            return None
        if choice == MENU_REFRESH:
            continue

        path = os.path.join(RESUME_FOLDER, choice)
        text, problem = load_resume(path)
        if problem:
            display_error(problem)
            continue

        display_message(f"Resume loaded ({len(text)} characters).")
        return {"resume_path": path, "resume_text": text}


def prompt_filters():
    """Ask for job preferences and return them as a dict."""
    display_message("\nYour job preferences (press Enter to accept the default)")

    while True:
        min_salary = prompt_number("Minimum monthly salary", 0, MAX_SALARY, default=0)
        max_salary = prompt_number("Maximum monthly salary", 0, MAX_SALARY, default=MAX_SALARY)
        if max_salary >= min_salary:
            break
        display_error("Maximum salary cannot be lower than minimum salary. Please enter both again.")

    max_years = prompt_number("Hide jobs that need more than how many years of experience",
                              0, MAX_EXPERIENCE_YEARS, default=MAX_EXPERIENCE_YEARS)
    job_type = prompt_choice("Job type", JOB_TYPES)
    work_arrangement = prompt_choice("Work arrangement", WORK_ARRANGEMENTS)

    return {
        "min_salary": min_salary,
        "max_salary": max_salary,
        "max_years_experience": max_years,
        "job_type": job_type,
        "work_arrangement": work_arrangement,
    }


def check_filters(values):
    """The web version of prompt_filters: check filters typed into a form (any dict-like of text).
    Return (filters, problem), with filters in the same shape as prompt_filters."""
    numbers = {}
    limits = {
        "min_salary": (0, MAX_SALARY),
        "max_salary": (MAX_SALARY, MAX_SALARY),
        "max_years_experience": (MAX_EXPERIENCE_YEARS, MAX_EXPERIENCE_YEARS),
    }
    for name, (default, maximum) in limits.items():
        text = values.get(name, "").strip().replace(",", "").replace("$", "")
        try:
            numbers[name] = int(text) if text else default
        except ValueError:
            return {}, "Salary and experience must be whole numbers."
        if not 0 <= numbers[name] <= maximum:
            return {}, f"Please keep {name.replace('_', ' ')} between 0 and {maximum}."
    if numbers["max_salary"] < numbers["min_salary"]:
        return {}, "Maximum salary cannot be lower than minimum salary."

    job_type = values.get("job_type", ANY)
    work_arrangement = values.get("work_arrangement", ANY)
    if job_type not in JOB_TYPES or work_arrangement not in WORK_ARRANGEMENTS:
        return {}, "Please pick a job type and work arrangement from the lists."
    return {**numbers, "job_type": job_type, "work_arrangement": work_arrangement}, ""


def collect_input(ai_name="an AI service"):
    """Collect the resume and filters for one search. ai_name is shown in the privacy notice."""
    display_header("NEW SEARCH")
    display_message(f"Note: your resume's text is sent to {ai_name} to analyse it. Remove anything you'd "
                    "rather not share (e.g. your home address or NRIC) first.")
    record = prompt_resume()
    if record is None:
        return None
    record["filters"] = prompt_filters()
    return record


def shorten(text, width):
    """Cut text to fit a column."""
    text = str(text)
    if len(text) <= width:
        return text
    return text[:width - 3] + "..."


def join_list(value):
    """Join a list into comma-separated text."""
    if not isinstance(value, list) or not value:
        return "None"
    names = []
    for item in value:
        if isinstance(item, dict):
            item = item.get("job_skill", "")
        names.append(str(item))
    return ", ".join(names)


def format_salary(job):
    """Return a job's salary range as text."""
    low = job.get("min_salary", 0)
    high = job.get("max_salary", 0)
    if not isinstance(low, (int, float)) or not isinstance(high, (int, float)) or (low == 0 and high == 0):
        return "Not stated"
    return f"${low:,.0f} - ${high:,.0f} {job.get('salary_period', 'monthly')}"


def tidy_title(title):
    """Job titles written in capitals, e.g. "SERVICE CREW" -> "Service Crew". Others are left as they are."""
    if not title.isupper():
        return title
    return " ".join(word if word in ACRONYMS else word.capitalize() for word in title.split())


def short_location(location):
    """First place in a MyCareersFuture district, e.g. "D19 Hougang, Sengkang, ..." -> "Hougang"."""
    place = re.sub(r"^D\d+\s+", "", location or "").split(",")[0].strip()
    return place or location


def experience_text(years):
    """A job's minimum experience as text, e.g. 0 -> "No experience needed", 2 -> "2+ years"."""
    if not years:
        return "No experience needed"
    return f"{years}+ year{'s' if years != 1 else ''}"


def same_skill(job_skill, candidate_skill):
    """True if two skill names are the same apart from case and spaces, e.g. "SQL" and "sql",
    so a match under another name ("Stakeholder Management" by "Stakeholder Communication") can be shown."""
    return job_skill.strip().lower() == candidate_skill.strip().lower()


def display_profile(profile):
    """Show what the AI found in the resume."""
    display_header("YOUR PROFILE")
    print(f"Name          : {profile.get('candidate_name', 'N/A')}")
    print(f"Qualification : {profile.get('qualification_detail', 'N/A')} "
          f"({profile.get('highest_qualification', 'N/A')})")
    print(f"Experience    : {profile.get('years_of_experience', 'N/A')} years full-time, "
          f"{profile.get('internship_months', 'N/A')} months internship")
    print(f"Level         : {profile.get('seniority_level', 'N/A')}")
    print(f"Skills        : {join_list(profile.get('core_skills'))}")
    print(f"Certifications: {join_list(profile.get('certifications'))}")


def display_record(job):
    """Show every detail of one job."""
    display_header(shorten(job.get("title", "N/A"), LINE_WIDTH))
    print(f"Company         : {job.get('company', 'N/A')}")
    print(f"Location        : {job.get('location', 'N/A')}")
    print(f"Salary          : {format_salary(job)}")
    print(f"Job type        : {join_list(job.get('employment_types'))}")
    print(f"Work arrangement: {job.get('work_arrangement', 'N/A')}")
    print(f"Experience      : {job.get('min_years_experience', 'N/A')} years minimum")
    print(f"Education       : {job.get('min_education', 'N/A')}")
    print(f"Skills you have : {join_list(job.get('matched_skills'))}")
    print(f"Skills missing  : {join_list(job.get('missing_skills'))}")
    print(f"Why             : {job.get('suitability_reason', 'N/A')}")
    print(f"Link            : {job.get('job_url', 'N/A')}")


def display_list(jobs):
    """Show a numbered table of jobs."""
    print(f"{'No.':<4} {'Title':<34} {'Company':<20} {'Salary'}")
    print("-" * LINE_WIDTH)
    for number, job in enumerate(jobs, start=1):
        print(f"{number:<4} {shorten(job.get('title', 'N/A'), 34):<34} "
              f"{shorten(job.get('company', 'N/A'), 20):<20} {format_salary(job)}")


def display_result(profile, jobs):
    """Show the profile and the matching jobs."""
    display_profile(profile)
    display_header(f"MATCHING JOBS ({len(jobs)})")
    if not jobs:
        display_message("No jobs matched. Try a new search with wider filters.")
        return
    display_list(jobs)


def check_ai_result(profile, jobs):
    """Return what is wrong with the AI result, or an empty string."""
    if profile is None:
        return ("The AI could not analyse your resume. Check your internet connection and API key, "
                "or wait a minute if the AI service is busy.")
    if not isinstance(profile, dict) or not isinstance(jobs, list):
        return "The AI returned data in an unexpected format."
    return ""


def prompt_missing_data(profile):
    """Ask for details the AI could not find in the resume."""
    if profile.get("highest_qualification", "None") == "None":
        display_message("\nThe AI could not find a completed qualification in your resume.")
        level = prompt_choice("Your highest completed qualification", EDUCATION_LEVELS)
        profile["highest_qualification"] = level
        if not profile.get("qualification_detail"):
            profile["qualification_detail"] = level

    if not profile.get("core_skills"):
        display_message("\nThe AI could not find any skills in your resume.")
        skills = []
        while not skills:
            answer = prompt_text("Your skills, separated by commas (e.g. Python, Excel)")
            for skill in answer.split(","):
                if skill.strip():
                    skills.append(skill.strip())
            if not skills:
                display_error("Please enter at least one skill.")
        profile["core_skills"] = skills

    return profile


def fill_missing_data(profile, qualification, skills_text):
    """The web version of prompt_missing_data: fill in what the AI could not find in the resume
    from the search form's optional fields. skills_text is comma-separated, e.g. "Python, Excel"."""
    if profile["highest_qualification"] == "None" and qualification in EDUCATION_LEVELS:
        profile["highest_qualification"] = qualification
        profile["qualification_detail"] = profile["qualification_detail"] or qualification
    skills = [skill.strip() for skill in skills_text.split(",") if skill.strip()]
    if not profile["core_skills"] and skills:
        profile["core_skills"] = skills
    return profile


def run_search(process_function, record):
    """Run one search, offering a retry if it fails."""
    while True:
        display_message("\nAnalysing your resume and searching for jobs. This can take a minute...")
        try:
            profile, jobs = process_function(record)
        except Exception as error:
            display_error(f"Something went wrong during the search ({type(error).__name__}: {error}).")
            profile, jobs = None, []

        problem = check_ai_result(profile, jobs)
        if problem == "":
            return profile, [job for job in jobs if isinstance(job, dict)]

        display_error(problem)
        if not prompt_yes_no("Try again?"):
            return None, []


# Progress of the web interface's background tasks (a search, or tailoring a resume), by id:
# {"status": "running" | "done" | "error" | "cancelled", "percent", "message", "result", "error"}
# (tailoring tasks also have "job_url"). The browser asks for this every second to draw the progress bar.
# Kept in memory only (lost on restart); the jobs themselves are saved by the data layer.
# Background threads update it, so every change goes through the lock.
tasks = {}
tasks_lock = threading.Lock()


def start_task(**extra):
    """Add a running task and return its id. extra: e.g. job_url for tailoring."""
    task_id = uuid.uuid4().hex
    with tasks_lock:
        tasks[task_id] = {"status": "running", "percent": 0, "message": "Starting...",
                          "result": None, "error": None, **extra}
    return task_id


def update_task(task_id, **changes):
    """Update a task's progress. Raises InterruptedError (built into Python) if the user
    cancelled it, so the background thread stops at its next update."""
    with tasks_lock:
        if tasks[task_id]["status"] == "cancelled":
            raise InterruptedError("cancelled by the user")
        tasks[task_id].update(changes)


def get_task(task_id):
    """A copy of one task, or None if there is no task with that id."""
    with tasks_lock:
        task = tasks.get(task_id)
        return dict(task) if task is not None else None


def cancel_task(task_id):
    """Mark a running task as cancelled; its thread stops at its next update_task."""
    with tasks_lock:
        task = tasks.get(task_id)
        if task is not None and task["status"] == "running":
            task["status"] = "cancelled"


def find_recent_job(job_url):
    """A job by its link from a finished search still in memory, with the resume that found it, or None."""
    with tasks_lock:
        results = [task["result"] for task in tasks.values()
                   if task["status"] == "done" and "all_jobs" in task["result"]]
    for result in results:
        for job in result["all_jobs"]:
            if job["job_url"] == job_url:
                return {**job, "resume": result["resume"]}
    return None
